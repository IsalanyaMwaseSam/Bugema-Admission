import base64
from django.shortcuts import get_object_or_404, render, redirect
from django.urls import reverse
from requests import HTTPError
from Application.models import *
from sendgrid import SendGridAPIClient, Attachment
from sendgrid.helpers.mail import Mail, Disposition, FileContent, FileName, FileType
from django.contrib import messages
from docx import Document
from docx.shared import Inches
from io import BytesIO
from django.conf import settings
import os
import configparser

config = configparser.ConfigParser()
config.read('config.py')

api_key = config.get('sendgrid', 'api_key')


# Create your views here.
def dashboard(request):
    pending_applications = Application.objects.filter(status='pending')
    approved_applications = Application.objects.filter(status='approved')
    rejected_applications = Application.objects.filter(status='rejected')
    context = {
        'approved_count': approved_applications.count(),
        'pending_count': pending_applications.count(),
        'rejected_count': rejected_applications.count(),
        'total_count': approved_applications.count() + pending_applications.count() + rejected_applications.count()
    }
    return render(request, 'Administration/dashboard.html', context)

def ApprovedApplications(request):
    approved_applications = Application.objects.filter(status='approved')
    pending_applications = Application.objects.filter(status='pending')
    rejected_applications = Application.objects.filter(status='rejected')
    total_approved = approved_applications.count()
    total_pending = pending_applications.count()
    total_rejected = rejected_applications.count()
    context = {
        'applications': approved_applications,
        'total_pending': total_pending,
        'total_approved': total_approved,
        'total_rejected': total_rejected,
      
    }
    return render(request, 'Administration/approved.html', context)

def RejectedApplications(request):
    pending_applications = Application.objects.filter(status='pending')
    approved_applications = Application.objects.filter(status='approved')
    rejected_applications = Application.objects.filter(status='rejected')
    total_approved = approved_applications.count()
    total_pending = pending_applications.count()
    total_rejected = rejected_applications.count()
    context = {
        'applications': rejected_applications,
        'total_pending': total_pending,
        'total_approved': total_approved,
        'total_rejected': total_rejected,
      
    }
    return render(request, 'Administration/rejected.html', context)

def PendingApplications(request):
    pending_applications = Application.objects.filter(status='pending')
    approved_applications = Application.objects.filter(status='approved')
    rejected_applications = Application.objects.filter(status='rejected')
    total_approved = approved_applications.count()
    total_pending = pending_applications.count()
    total_rejected = rejected_applications.count()
    context = {
        'applications': pending_applications,
        'total_pending': total_pending,
        'total_approved': total_approved,
        'total_rejected': total_rejected,
      
    }
    return render(request, 'Administration/pending.html', context)

def approve_application(request, application_id):
    application = get_object_or_404(Application, id=application_id)
    application.status = 'approved'
    application.save()

    template_path = 'static/documents/Rugaaju Wicliff_ Admission0001 _1_.docx'
    doc = Document(template_path)

    # Replace placeholders with applicant's name and program
    name = application.contact.name
    program = application.program.name
    for p in doc.paragraphs:
        if '{{name}}' in p.text:
            p.text = p.text.replace('{{name}}', name)
        if '{{program}}' in p.text:
            p.text = p.text.replace('{{program}}', program)

    # Save the modified document as a new file
    output_path = os.path.join(settings.MEDIA_ROOT, 'uploads', f'admission_letter_{application.pk}.docx')
    doc.save(output_path)

    # attachment = Attachment()

        # attachment.content = f.read()
        # attachment.type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        # attachment.filename = f'admission_letter_{application.pk}.docx'
        # attachment.disposition = Disposition('attachment')

    message = Mail(
        from_email='admissions@bugemauniv.ac.ug',
        to_emails=application.contact.email,
        subject='Congratulations! Your application has been approved',
        html_content=f'<p>Dear {name},</p><p>We are pleased to inform you that your application for {program} has been approved.</p><p>Best regards,<br>Admissions office</p>'
    )
    with open(output_path, 'rb') as f:
        data = f.read()
        f.close()
    encoded_file = base64.b64encode(data).decode()

    attachedFile = Attachment(
        FileContent(encoded_file),
        FileName('admission_letter_{application.pk}.docx'),
        FileType('application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
        Disposition('attachment')
    )
    message.attachment = attachedFile

    try:
        sg = SendGridAPIClient(api_key)
        response = sg.send(message)
        print(response.status_code)
        print(response.body)
        print(response.headers)
    except Exception as e:
        print(f'Other error occurred: {e}')
        print(str(e))
        messages.error(request, 'Failed to send email')
    
    return redirect('approved-applications')

def reject_application(request, application_id):
    application = get_object_or_404(Application, id=application_id)
    application.status = 'rejected'
    application.save()
    return redirect(reverse('rejected-applications'))

